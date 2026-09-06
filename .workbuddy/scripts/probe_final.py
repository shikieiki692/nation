# -*- coding: utf-8 -*-
"""最终探查：页脚真实内容 / 正文颜色归属 / 文档开头结构"""
import sys, io, zipfile, re
from collections import Counter, defaultdict
from pathlib import Path
import docx

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
ROOT = Path(r"C:\Obsidion\妙妙屋\00-首页\题组Word\习题书\学生版")

f1 = ROOT / "第一篇-化学原理" / "1-热力学.docx"

print("### 页脚 footer1.xml 全文（去命名空间前缀）")
z = zipfile.ZipFile(str(f1))
x = z.read('word/footer1.xml').decode('utf-8', 'replace')
x = re.sub(r'\s+xmlns:[a-z0-9]+="[^"]*"', '', x)
print(re.sub(r'>\s*<', '>\n<', x))

print("\n### 正文 w:color 归属（仅 document.xml）")
agg = defaultdict(Counter)
for f in sorted(ROOT.rglob('*.docx')):
    z = zipfile.ZipFile(str(f))
    xml = z.read('word/document.xml').decode('utf-8', 'replace')
    # 对每个段落，找 pStyle 和内部 color
    for pm in re.finditer(r'<w:p>(.*?)</w:p>', xml, re.S):
        blk = pm.group(1)
        ps = re.search(r'<w:pStyle w:val="([^"]+)"', blk)
        style = ps.group(1) if ps else 'Normal'
        for cm in re.finditer(r'<w:color w:val="([^"]+)"', blk):
            agg[cm.group(1)][style] += 1
for col, styles in sorted(agg.items(), key=lambda kv: -sum(kv[1].values())):
    print(f"  {col}: 共{sum(styles.values())} -> {dict(styles)}")

print("\n### 1-热力学.docx 前 30 段（样式/字号/文本）")
d = docx.Document(str(f1))
for i, p in enumerate(d.paragraphs[:30]):
    r0 = p.runs[0] if p.runs else None
    sz = r0.font.size.pt if r0 and r0.font.size else None
    b = r0.bold if r0 else None
    print(f"  [{i:>2}] {p.style.name:<18} {str(sz):>5} b={str(b):<5} | {p.text[:56]!r}")

print("\n### 图片尺寸分布（前 15）")
z2 = zipfile.ZipFile(str(ROOT / "第三篇-有机化学" / "1-结构基础与波谱分析.docx"))
xml2 = z2.read('word/document.xml').decode('utf-8', 'replace')
sizes = re.findall(r'<wp:extent cx="(\d+)" cy="(\d+)"', xml2)
for cx, cy in sizes[:15]:
    print(f"  {int(cx)/360000:.2f}cm x {int(cy)/360000:.2f}cm")

# 全库图片尺寸
print("\n### 全库图片尺寸统计")
allsz = []
for f in sorted(ROOT.rglob('*.docx')):
    z = zipfile.ZipFile(str(f))
    try:
        xml = z.read('word/document.xml').decode('utf-8', 'replace')
    except KeyError:
        continue
    for cx, cy in re.findall(r'<wp:extent cx="(\d+)" cy="(\d+)"', xml):
        allsz.append((int(cx) / 360000, int(cy) / 360000))
if allsz:
    ws = sorted(w for w, h in allsz)
    print(f"  数量 {len(allsz)} | 宽度 min={ws[0]:.2f} 中位={ws[len(ws)//2]:.2f} max={ws[-1]:.2f}")
    print(f"  超宽(>16.6cm 即超出A4版心) 数量: {sum(1 for w,h in allsz if w>16.6)}")
