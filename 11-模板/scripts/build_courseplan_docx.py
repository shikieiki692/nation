# -*- coding: utf-8 -*-
"""把课程计划 md 转成与原样板同版式的 docx（通用表格版）。

版式（沿用原样板，逐项提取自 C:\\Users\\蕾赛\\Downloads\\初三未央化学竞赛课程计划..docx）：
  - 页面  A4(11906x16838)，页边距 上下 1440 / 左右 1800，docGrid lines linePitch 312
  - 字体  西文 Times New Roman，中文 宋体
  - 字号  标题 14pt（居中），章节标题与表格正文 一律 10pt，全程不加粗
  - 段落  段后 0，行距 240(auto=单倍)
  - 表格  框线 single sz=4；单元格垂直居中、文本水平居中
  - 单元格内以「｜」分行的内容拆成多个段落

v2（2026-09-21）：表格改为**通用渲染**——表头一律保留、列宽按内容自适应；
仅两类表做纵向合并：模块表（章＋次数 合并、教材来源 合并）与章后课表（章 合并）。
"""
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

SRC = Path(r"C:\Obsidion\妙妙屋\备课思路\未央化学竞赛课程计划（初二至高二）.md")
OUT = Path(r"C:\Obsidion\妙妙屋\备课思路\未央化学竞赛课程计划（初二至高二）.docx")

SEP = "｜"
EAST, WEST = "宋体", "Times New Roman"
TOTAL_W = 8391            # 与旧版一致
W_MIN = 620

MODULE_HEADER = ["章", "节", "知识点", "教材来源", "次数"]
REVIEW_HEADER = ["章", "课型", "知识点", "课次"]


# ---------------- 底层工具 ----------------

def clean_md_text(s):
    s = re.sub(r"\[\[([^\]|]+)\|([^\]]+)\]\]", r"\2", s)
    s = re.sub(r"\[\[([^\]]+)\]\]", lambda m: m.group(1).split("/")[-1], s)
    s = re.sub(r"\*\*([^*]+)\*\*", r"\1", s)
    s = re.sub(r"`([^`]+)`", r"\1", s)
    s = re.sub(r"^>\s*", "", s)
    s = re.sub(r"^[-*]\s+", "· ", s)
    return s.strip()


def disp_len(s):
    """CJK 记 2、其余记 1 的显示宽度。"""
    n = 0
    for ch in s:
        n += 2 if ord(ch) > 0x2E80 else 1
    return n


def set_run(run, hp):
    run.font.size = Pt(hp / 2)
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    rf.set(qn("w:ascii"), WEST)
    rf.set(qn("w:hAnsi"), WEST)
    rf.set(qn("w:cs"), WEST)
    rf.set(qn("w:eastAsia"), EAST)
    for tag in ("w:sz", "w:szCs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = OxmlElement(tag)
            rPr.append(e)
        e.set(qn("w:val"), str(hp))


def style_para(p, align=None, hp=20):
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    if align is not None:
        p.alignment = align
    for r in p.runs:
        set_run(r, hp)
    return p


def add_block(doc, text, align=None, hp=20):
    p = doc.add_paragraph()
    p.add_run(text)
    return style_para(p, align, hp)


def table_borders(tbl):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    b = OxmlElement("w:tblBorders")
    for tag in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + tag)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "auto")
        b.append(e)
    tblPr.append(b)


def set_grid(tbl, widths):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn("w:tblW")):
        tblPr.remove(old)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(sum(widths)))
    w.set(qn("w:type"), "dxa")
    tblPr.append(w)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    old = tbl._tbl.find(qn("w:tblGrid"))
    if old is not None:
        tbl._tbl.remove(old)
    grid = OxmlElement("w:tblGrid")
    for ww in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(ww))
        grid.append(gc)
    tblPr.addnext(grid)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False


def set_vmerge(cell, val=None):
    tcPr = cell._tc.get_or_add_tcPr()
    vm = tcPr.find(qn("w:vMerge"))
    if vm is None:
        vm = OxmlElement("w:vMerge")
        ref = tcPr.find(qn("w:tcW"))
        if ref is None:
            ref = tcPr.find(qn("w:gridSpan"))
        if ref is not None:
            ref.addnext(vm)
        else:
            tcPr.insert(0, vm)
    if val is not None:
        vm.set(qn("w:val"), val)


def set_cell(cell, lines, hp=20):
    if not lines:
        lines = [""]
    for i, it in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.add_run(it)
        style_para(p, WD_ALIGN_PARAGRAPH.CENTER, hp)


def calc_widths(rows, ncol):
    """列宽 = 「短标签列给足、长文本列分享剩余」。

    10pt 汉字宽约 200 twips，disp_len 以「半角」为单位（汉字=2），故 1 单位 ≈ 100 twips；
    单元格左右内边距合计约 220 twips。need ≤ SMALL 的短标签列按需给足，避免折行；
    其余列按 need 比例分享剩余宽度。
    """
    UNIT, PAD, HARD_MIN, CAP = 100, 220, 560, int(TOTAL_W * 0.40)
    SMALL = int(TOTAL_W * 0.28)
    need = []
    for c in range(ncol):
        mx = 2
        for r in rows:
            v = max((disp_len(x) for x in r[c].split(SEP)), default=0)
            mx = max(mx, v)
        need.append(min(CAP, mx * UNIT + PAD))

    small = [i for i, n in enumerate(need) if n <= SMALL]
    big = [i for i, n in enumerate(need) if n > SMALL]
    fixed = sum(need[i] for i in small)
    w = [0] * ncol
    if not big or fixed + HARD_MIN * len(big) > TOTAL_W:
        # 保护不了短标签列（列太挤）→ 全表按比例缩放
        k = TOTAL_W / sum(need)
        w = [max(HARD_MIN, int(n * k)) for n in need]
    else:
        rest = TOTAL_W - fixed
        tot_big = sum(need[i] for i in big)
        for i in small:
            w[i] = need[i]
        for i in big:
            w[i] = max(HARD_MIN, int(rest * need[i] / tot_big))
    w[w.index(max(w))] += TOTAL_W - sum(w)
    return w


# ---------------- 解析 md ----------------

lines = SRC.read_text(encoding="utf-8").split("\n")
title = None
blocks = []
i = 0
while i < len(lines):
    s = lines[i].strip()
    if s.startswith("# ") and not s.startswith("## "):
        title = s[2:].strip()
        i += 1
        continue
    if s.startswith("## "):
        blocks.append(("h2", s[3:].strip(), None))
        i += 1
        continue
    if s.startswith("### "):
        blocks.append(("h3", s[4:].strip(), None))
        i += 1
        continue
    if s.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
            i += 1
        rows = [r for r in rows if not (set("".join(r)) <= set(":- ") and r)]
        blocks.append(("table", None, rows))
        continue
    if title is not None and s:
        blocks.append(("text", s, None))
    i += 1

assert title, "未找到 H1 标题"

# ---------------- 生成 docx ----------------

doc = Document()
sec = doc.sections[0]
sec.page_width = Twips(11906)
sec.page_height = Twips(16838)
sec.top_margin = Twips(1440)
sec.bottom_margin = Twips(1440)
sec.left_margin = Twips(1800)
sec.right_margin = Twips(1800)

normal = doc.styles["Normal"]
normal.font.size = Pt(10)
normal.font.name = WEST
rpr = normal.element.get_or_add_rPr()
rf = rpr.find(qn("w:rFonts"))
if rf is None:
    rf = OxmlElement("w:rFonts")
    rpr.insert(0, rf)
for k, v in (("w:ascii", WEST), ("w:hAnsi", WEST), ("w:cs", WEST), ("w:eastAsia", EAST)):
    rf.set(qn(k), v)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.line_spacing = 1.0

sect_pr = sec._sectPr
dg = sect_pr.find(qn("w:docGrid"))
if dg is None:
    dg = OxmlElement("w:docGrid")
    sect_pr.append(dg)
dg.set(qn("w:type"), "lines")
dg.set(qn("w:linePitch"), "312")

add_block(doc, title, WD_ALIGN_PARAGRAPH.CENTER, hp=28)

stats = []
n_merge = 0
n_mod = 0
for kind, name, rows in blocks:
    if kind in ("h2", "h3", "text"):
        add_block(doc, clean_md_text(name), None, hp=20)
        continue

    if not rows:
        continue
    ncol = max(len(r) for r in rows)
    rows = [(list(r) + [""] * ncol)[:ncol] for r in rows]
    hdr = rows[0]
    if hdr == MODULE_HEADER:
        n_mod += 1
    widths = calc_widths(rows, ncol)
    tbl = doc.add_table(rows=len(rows), cols=ncol)
    set_grid(tbl, widths)
    table_borders(tbl)
    for ri in range(len(rows)):
        trPr = tbl.rows[ri]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))   # 整行不跨页
        for ci in range(ncol):
            items = [clean_md_text(x) for x in rows[ri][ci].split(SEP) if x.strip()]
            set_cell(tbl.cell(ri, ci), items or [""])
            cell = tbl.cell(ri, ci)
            cell.width = Twips(widths[ci])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    if hdr == MODULE_HEADER:
        gs = [ri for ri in range(len(rows)) if rows[ri][0].strip()] + [len(rows)]
        for a, b in zip(gs, gs[1:]):
            if b - a > 1:
                # 章(0) / 教材来源(3) / 次数(4) 均为章级字段 → 一并纵向合并
                for col in (0, 3, 4):
                    set_vmerge(tbl.cell(a, col), "restart")
                    for k in range(a + 1, b):
                        set_vmerge(tbl.cell(k, col), None)
                n_merge += 3
    elif hdr == REVIEW_HEADER:
        gs = [ri for ri in range(len(rows)) if rows[ri][0].strip()] + [len(rows)]
        for a, b in zip(gs, gs[1:]):
            if b - a > 1:
                set_vmerge(tbl.cell(a, 0), "restart")
                for k in range(a + 1, b):
                    set_vmerge(tbl.cell(k, 0), None)
                n_merge += 1

    stats.append((len(rows), ncol))

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUT))
print("[ok]", OUT)
print("[size]", OUT.stat().st_size, "bytes")
print("[tables]", len(stats), " 模块表:", n_mod, " 纵向合并起点:", n_merge)
for n, c in stats:
    print("   行=%-3d 列=%d" % (n, c))
